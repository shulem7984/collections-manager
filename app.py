import os, io, csv, sqlite3, secrets, shutil
from datetime import datetime, date
from functools import wraps
from pathlib import Path
from flask import Flask, g, render_template, request, redirect, url_for, session, flash, send_file, abort
from werkzeug.security import generate_password_hash, check_password_hash

BASE = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get('COLLECTIONS_DB', BASE / 'collections.db'))
BACKUP_DIR = Path(os.environ.get('COLLECTIONS_BACKUP_DIR', BASE / 'backups'))
BACKUP_DIR.mkdir(parents=True, exist_ok=True)
app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY') or secrets.token_hex(32)
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax', SESSION_COOKIE_SECURE=os.environ.get('COOKIE_SECURE', '0') == '1', MAX_CONTENT_LENGTH=12 * 1024 * 1024)

SCHEMA = '''
CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, username TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL, role TEXT NOT NULL DEFAULT 'employee', active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS companies(id INTEGER PRIMARY KEY, name TEXT NOT NULL, contact_name TEXT DEFAULT '', email TEXT DEFAULT '', phone TEXT DEFAULT '', commission_rate REAL NOT NULL DEFAULT 0, notes TEXT DEFAULT '', created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS accounts(id INTEGER PRIMARY KEY, company_id INTEGER NOT NULL REFERENCES companies(id) ON DELETE CASCADE, account_number TEXT DEFAULT '', customer_name TEXT NOT NULL, email TEXT DEFAULT '', phone TEXT DEFAULT '', address TEXT DEFAULT '', original_balance REAL NOT NULL DEFAULT 0, balance REAL NOT NULL DEFAULT 0, status TEXT NOT NULL DEFAULT 'Open', assigned_to INTEGER REFERENCES users(id), next_followup TEXT DEFAULT '', created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS account_notes(id INTEGER PRIMARY KEY, account_id INTEGER NOT NULL REFERENCES accounts(id) ON DELETE CASCADE, user_id INTEGER REFERENCES users(id), note_type TEXT NOT NULL DEFAULT 'Note', body TEXT NOT NULL, created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS collections(id INTEGER PRIMARY KEY, account_id INTEGER NOT NULL REFERENCES accounts(id) ON DELETE CASCADE, amount REAL NOT NULL CHECK(amount > 0), collected_on TEXT NOT NULL, method TEXT DEFAULT 'Other', reference TEXT DEFAULT '', entered_by INTEGER REFERENCES users(id), created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS audit_log(id INTEGER PRIMARY KEY, user_id INTEGER REFERENCES users(id), action TEXT NOT NULL, details TEXT DEFAULT '', created_at TEXT NOT NULL);
'''

def now(): return datetime.now().isoformat(timespec='seconds')
def db():
    if 'db' not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute('PRAGMA foreign_keys = ON')
    return g.db
@app.teardown_appcontext
def close_db(exc):
    conn = g.pop('db', None)
    if conn is not None: conn.close()

def init_db():
    conn = sqlite3.connect(DB_PATH)
    conn.executescript(SCHEMA)
    conn.commit()
    if not conn.execute('SELECT id FROM users LIMIT 1').fetchone():
        username = os.environ.get('ADMIN_USERNAME', 'admin')
        password = os.environ.get('ADMIN_PASSWORD', 'ChangeMe-Immediately-2026!')
        conn.execute('INSERT INTO users(username,password_hash,role,created_at) VALUES(?,?,?,?)', (username, generate_password_hash(password), 'admin', now()))
        conn.commit()
    conn.close()

@app.before_request
def csrf_setup():
    if 'csrf' not in session: session['csrf'] = secrets.token_urlsafe(24)
    if request.method == 'POST':
        token = request.form.get('_csrf', '')
        if not token or not secrets.compare_digest(token, session.get('csrf', '')): abort(400, 'Invalid security token. Refresh the page and try again.')

def csrf(): return session.get('csrf', '')
app.jinja_env.globals['csrf_token'] = csrf
app.jinja_env.globals['today'] = lambda: date.today().isoformat()

def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get('user_id'): return redirect(url_for('login', next=request.path))
        user = db().execute('SELECT * FROM users WHERE id=? AND active=1', (session['user_id'],)).fetchone()
        if not user: session.clear(); return redirect(url_for('login'))
        g.user = user
        return fn(*args, **kwargs)
    return wrapped

def admin_required(fn):
    @wraps(fn)
    @login_required
    def wrapped(*args, **kwargs):
        if g.user['role'] != 'admin': abort(403)
        return fn(*args, **kwargs)
    return wrapped

def account_visible(account_id, user=None):
    user = user or g.user
    row = db().execute('SELECT assigned_to FROM accounts WHERE id=?', (account_id,)).fetchone()
    return bool(row and (user['role'] == 'admin' or row['assigned_to'] == user['id']))

def require_account_access(account_id):
    if not account_visible(account_id):
        # Do not reveal whether an inaccessible account exists.
        abort(404)

def audit(action, details=''):
    db().execute('INSERT INTO audit_log(user_id,action,details,created_at) VALUES(?,?,?,?)', (session.get('user_id'), action, details[:500], now()))
    db().commit()

def money(v): return '${:,.2f}'.format(float(v or 0))
app.jinja_env.filters['money'] = money

@app.route('/login', methods=['GET','POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username','').strip()
        row = db().execute('SELECT * FROM users WHERE username=? AND active=1', (username,)).fetchone()
        if row and check_password_hash(row['password_hash'], request.form.get('password','')):
            session.clear(); session['user_id'] = row['id']; session['csrf'] = secrets.token_urlsafe(24)
            db().execute('INSERT INTO audit_log(user_id,action,details,created_at) VALUES(?,?,?,?)', (row['id'],'login','Successful login',now())); db().commit()
            return redirect(url_for('dashboard'))
        flash('Invalid username or password.', 'error')
    return render_template('login.html')
@app.route('/logout', methods=['POST'])
def logout():
    session.clear(); return redirect(url_for('login'))

@app.route('/')
@login_required
def dashboard():
    conn=db()
    scope = '' if g.user['role'] == 'admin' else ' AND a.assigned_to=?'
    scope_args = [] if g.user['role'] == 'admin' else [g.user['id']]
    stats={
      'companies':conn.execute('SELECT COUNT(DISTINCT company_id) n FROM accounts a WHERE 1=1'+scope, scope_args).fetchone()['n'] if g.user['role'] != 'admin' else conn.execute('SELECT COUNT(*) n FROM companies').fetchone()['n'],
      'accounts':conn.execute('SELECT COUNT(*) n FROM accounts a WHERE 1=1'+scope, scope_args).fetchone()['n'],
      'outstanding':conn.execute("SELECT COALESCE(SUM(balance),0) n FROM accounts a WHERE status!='Closed'"+scope, scope_args).fetchone()['n'],
      'collected':conn.execute('SELECT COALESCE(SUM(c.amount),0) n FROM collections c JOIN accounts a ON a.id=c.account_id WHERE 1=1'+scope, scope_args).fetchone()['n'],
      'commission':conn.execute('SELECT COALESCE(SUM(c.amount * co.commission_rate / 100.0),0) n FROM collections c JOIN accounts a ON a.id=c.account_id JOIN companies co ON co.id=a.company_id WHERE 1=1'+scope, scope_args).fetchone()['n']}
    due=conn.execute('SELECT a.*,c.name company_name FROM accounts a JOIN companies c ON c.id=a.company_id WHERE a.next_followup<>"" AND a.next_followup<=? AND a.status<>"Closed"'+scope+' ORDER BY a.next_followup LIMIT 10',[date.today().isoformat()]+scope_args).fetchall()
    recent=conn.execute('SELECT c.*,a.customer_name,a.account_number,co.name company_name FROM collections c JOIN accounts a ON a.id=c.account_id JOIN companies co ON co.id=a.company_id WHERE 1=1'+scope+' ORDER BY c.collected_on DESC,c.id DESC LIMIT 8',scope_args).fetchall()
    return render_template('dashboard.html',stats=stats,due=due,recent=recent)

@app.route('/companies', methods=['GET','POST'])
@login_required
def companies():
    conn=db()
    if request.method=='POST':
        if g.user['role'] != 'admin': abort(403)
        name=request.form.get('name','').strip()
        rate=float(request.form.get('commission_rate') or 0)
        if not name: flash('Company name is required.','error')
        elif rate < 0 or rate > 100: flash('Commission rate must be between 0 and 100 percent.','error')
        else:
            conn.execute('INSERT INTO companies(name,contact_name,email,phone,commission_rate,notes,created_at) VALUES(?,?,?,?,?,?,?)',(name,request.form.get('contact_name','').strip(),request.form.get('email','').strip(),request.form.get('phone','').strip(),rate,request.form.get('notes','').strip(),now())); conn.commit(); audit('company_created',name); flash('Client company added.','success')
        return redirect(url_for('companies'))
    if g.user['role'] == 'admin':
        rows=conn.execute('SELECT c.*,COUNT(a.id) account_count,COALESCE(SUM(a.balance),0) outstanding FROM companies c LEFT JOIN accounts a ON a.company_id=c.id GROUP BY c.id ORDER BY c.name').fetchall()
    else:
        rows=conn.execute('SELECT c.*,COUNT(a.id) account_count,COALESCE(SUM(a.balance),0) outstanding FROM companies c JOIN accounts a ON a.company_id=c.id AND a.assigned_to=? GROUP BY c.id ORDER BY c.name',(g.user['id'],)).fetchall()
    return render_template('companies.html',companies=rows)
@app.route('/companies/<int:company_id>', methods=['GET','POST'])
@login_required
def company_detail(company_id):
    conn=db(); company=conn.execute('SELECT * FROM companies WHERE id=?',(company_id,)).fetchone()
    if not company: abort(404)
    if request.method=='POST':
        if g.user['role'] != 'admin': abort(403)
        rate=float(request.form.get('commission_rate') or 0)
        if rate < 0 or rate > 100: flash('Commission rate must be between 0 and 100 percent.','error'); return redirect(url_for('company_detail',company_id=company_id))
        conn.execute('UPDATE companies SET name=?,contact_name=?,email=?,phone=?,commission_rate=?,notes=? WHERE id=?',(request.form.get('name','').strip(),request.form.get('contact_name','').strip(),request.form.get('email','').strip(),request.form.get('phone','').strip(),rate,request.form.get('notes','').strip(),company_id)); conn.commit(); audit('company_updated',company['name']); flash('Company updated.','success'); return redirect(url_for('company_detail',company_id=company_id))
    if g.user['role'] == 'admin':
        accounts=conn.execute('SELECT a.*,u.username assigned_name FROM accounts a LEFT JOIN users u ON u.id=a.assigned_to WHERE a.company_id=? ORDER BY a.customer_name',(company_id,)).fetchall()
    else:
        accounts=conn.execute('SELECT a.*,u.username assigned_name FROM accounts a LEFT JOIN users u ON u.id=a.assigned_to WHERE a.company_id=? AND a.assigned_to=? ORDER BY a.customer_name',(company_id,g.user['id'])).fetchall()
    return render_template('company_detail.html',company=company,accounts=accounts,users=conn.execute('SELECT id,username FROM users WHERE active=1 ORDER BY username').fetchall())

@app.route('/accounts')
@login_required
def accounts():
    q=request.args.get('q','').strip(); status=request.args.get('status','')
    sql='SELECT a.*,c.name company_name,u.username assigned_name FROM accounts a JOIN companies c ON c.id=a.company_id LEFT JOIN users u ON u.id=a.assigned_to WHERE 1=1'; args=[]
    if g.user['role'] != 'admin': sql+=' AND a.assigned_to=?'; args.append(g.user['id'])
    if q:
        like='%'+q+'%'; sql+=' AND (a.customer_name LIKE ? OR a.account_number LIKE ? OR a.phone LIKE ? OR a.email LIKE ? OR c.name LIKE ?)'; args += [like]*5
    if status: sql+=' AND a.status=?'; args.append(status)
    sql+=' ORDER BY a.customer_name LIMIT 500'
    return render_template('accounts.html',accounts=db().execute(sql,args).fetchall(),q=q,status=status,companies=db().execute('SELECT id,name FROM companies ORDER BY name').fetchall(),users=db().execute('SELECT id,username FROM users WHERE active=1 ORDER BY username').fetchall())
@app.route('/accounts/add',methods=['POST'])
@login_required
def account_add():
    try:
        conn=db(); bal=float(request.form.get('balance') or 0); orig=float(request.form.get('original_balance') or bal)
        assigned = (int(request.form['assigned_to']) if request.form.get('assigned_to') else None) if g.user['role'] == 'admin' else g.user['id']
        conn.execute('INSERT INTO accounts(company_id,account_number,customer_name,email,phone,address,original_balance,balance,status,assigned_to,next_followup,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(int(request.form['company_id']),request.form.get('account_number','').strip(),request.form['customer_name'].strip(),request.form.get('email','').strip(),request.form.get('phone','').strip(),request.form.get('address','').strip(),orig,bal,request.form.get('status','Open'),assigned,request.form.get('next_followup',''),now())); conn.commit(); audit('account_created',request.form['customer_name']); flash('Customer account added.','success')
    except (ValueError,KeyError): flash('Please check the account details and amounts.','error')
    return redirect(request.referrer or url_for('accounts'))
@app.route('/accounts/<int:account_id>',methods=['GET','POST'])
@login_required
def account_detail(account_id):
    conn=db(); a=conn.execute('SELECT a.*,c.name company_name,c.commission_rate,u.username assigned_name FROM accounts a JOIN companies c ON c.id=a.company_id LEFT JOIN users u ON u.id=a.assigned_to WHERE a.id=?',(account_id,)).fetchone()
    if not a: abort(404)
    require_account_access(account_id)
    if request.method=='POST':
        new_assignee = (int(request.form['assigned_to']) if request.form.get('assigned_to') else None) if g.user['role'] == 'admin' else a['assigned_to']
        conn.execute('UPDATE accounts SET customer_name=?,account_number=?,email=?,phone=?,address=?,status=?,assigned_to=?,next_followup=? WHERE id=?',(request.form.get('customer_name','').strip(),request.form.get('account_number','').strip(),request.form.get('email','').strip(),request.form.get('phone','').strip(),request.form.get('address','').strip(),request.form.get('status','Open'),new_assignee,request.form.get('next_followup',''),account_id)); conn.commit(); audit('account_updated',a['customer_name']); flash('Account updated.','success'); return redirect(url_for('account_detail',account_id=account_id))
    notes=conn.execute('SELECT n.*,u.username FROM account_notes n LEFT JOIN users u ON u.id=n.user_id WHERE n.account_id=? ORDER BY n.id DESC',(account_id,)).fetchall()
    payments=conn.execute('SELECT c.*,u.username FROM collections c LEFT JOIN users u ON u.id=c.entered_by WHERE c.account_id=? ORDER BY c.collected_on DESC,c.id DESC',(account_id,)).fetchall()
    return render_template('account_detail.html',a=a,notes=notes,collections=payments,users=conn.execute('SELECT id,username FROM users WHERE active=1 ORDER BY username').fetchall())
@app.route('/accounts/<int:account_id>/notes',methods=['POST'])
@login_required
def add_note(account_id):
    require_account_access(account_id)
    body=request.form.get('body','').strip()
    if body:
        db().execute('INSERT INTO account_notes(account_id,user_id,note_type,body,created_at) VALUES(?,?,?,?,?)',(account_id,session['user_id'],request.form.get('note_type','Note'),body,now())); db().commit(); audit('note_added',f'Account {account_id}')
    flash('History entry added.','success'); return redirect(url_for('account_detail',account_id=account_id))
@app.route('/accounts/<int:account_id>/collect',methods=['POST'])
@login_required
def add_collection(account_id):
    require_account_access(account_id)
    conn=db(); a=conn.execute('SELECT * FROM accounts WHERE id=?',(account_id,)).fetchone()
    try:
        amount=round(float(request.form.get('amount','0')),2)
        if not a or amount<=0 or amount>a['balance']: raise ValueError()
        conn.execute('INSERT INTO collections(account_id,amount,collected_on,method,reference,entered_by,created_at) VALUES(?,?,?,?,?,?,?)',(account_id,amount,request.form.get('collected_on') or date.today().isoformat(),request.form.get('method','Other'),request.form.get('reference','').strip(),session['user_id'],now()))
        newbal=round(a['balance']-amount,2); status='Closed' if newbal<=0 else a['status']; conn.execute('UPDATE accounts SET balance=?,status=? WHERE id=?',(max(0,newbal),status,account_id)); conn.execute('INSERT INTO account_notes(account_id,user_id,note_type,body,created_at) VALUES(?,?,?,?,?)',(account_id,session['user_id'],'Collection',f'Collection recorded: {money(amount)}. Remaining balance: {money(max(0,newbal))}.',now())); conn.commit(); audit('collection_recorded',f'Account {account_id}: {amount}'); flash('Collection recorded and balance updated.','success')
    except (ValueError,TypeError): flash('Enter a valid collection amount no greater than the current balance.','error')
    return redirect(url_for('account_detail',account_id=account_id))

@app.route('/import',methods=['GET','POST'])
@admin_required
def import_csv():
    if request.method=='POST':
        file=request.files.get('file'); company_id=request.form.get('company_id')
        if not file or not file.filename: flash('Choose a CSV or Excel file.','error'); return redirect(url_for('import_csv'))
        if not company_id: flash('Choose a client company.','error'); return redirect(url_for('import_csv'))
        try:
            filename=(file.filename or '').lower()
            if filename.endswith('.xlsx'):
                from openpyxl import load_workbook
                workbook=load_workbook(file.stream, read_only=True, data_only=True)
                sheet=workbook.active
                iterator=sheet.iter_rows(values_only=True)
                first=next(iterator, None)
                fieldnames=[str(v).strip() if v is not None else '' for v in (first or [])]
                reader=[dict(zip(fieldnames, [str(v).strip() if v is not None else '' for v in row])) for row in iterator]
            else:
                raw=file.stream.read().decode('utf-8-sig')
                csv_reader=csv.DictReader(io.StringIO(raw))
                fieldnames=csv_reader.fieldnames or []
                reader=list(csv_reader)
            headers={h.strip().lower():h for h in fieldnames if h}
            def get(row,*names):
                for n in names:
                    if n in headers: return str(row.get(headers[n]) or '').strip()
                return ''
            count=0; conn=db()
            for row in reader:
                name=get(row,'customer_name','customer','name','debtor name')
                if not name: continue
                def num(v):
                    try:return float(v.replace('$','').replace(',','').strip() or 0)
                    except:return 0.0
                bal=num(get(row,'balance','current_balance','amount_due','outstanding_balance')); orig=num(get(row,'original_balance','original_amount','amount')) or bal
                conn.execute('INSERT INTO accounts(company_id,account_number,customer_name,email,phone,address,original_balance,balance,status,next_followup,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)',(int(company_id),get(row,'account_number','account #','account'),name,get(row,'email'),get(row,'phone','telephone'),get(row,'address'),orig,bal,get(row,'status') or 'Open',get(row,'next_followup','follow_up_date'),now())); count+=1
            conn.commit(); audit('csv_import',f'{count} accounts for company {company_id}'); flash(f'Imported {count} customer accounts.','success')
        except Exception: flash('Import failed. Check that the CSV or XLSX file has a customer/name column and valid balance values.','error')
        return redirect(url_for('accounts'))
    return render_template('import.html',companies=db().execute('SELECT id,name FROM companies ORDER BY name').fetchall())
@app.route('/reports')
@login_required
def reports():
    start=request.args.get('start',date.today().replace(day=1).isoformat()); end=request.args.get('end',date.today().isoformat())
    if g.user['role'] == 'admin':
        rows=db().execute('SELECT co.name company_name,COUNT(DISTINCT a.id) accounts,COALESCE(SUM(c.amount),0) collected,co.commission_rate FROM companies co LEFT JOIN accounts a ON a.company_id=co.id LEFT JOIN collections c ON c.account_id=a.id AND c.collected_on BETWEEN ? AND ? GROUP BY co.id ORDER BY co.name',(start,end)).fetchall()
    else:
        rows=db().execute('SELECT co.name company_name,COUNT(DISTINCT a.id) accounts,COALESCE(SUM(c.amount),0) collected,co.commission_rate FROM companies co JOIN accounts a ON a.company_id=co.id AND a.assigned_to=? LEFT JOIN collections c ON c.account_id=a.id AND c.collected_on BETWEEN ? AND ? GROUP BY co.id ORDER BY co.name',(g.user['id'],start,end)).fetchall()
    # Use separate query for outstanding so collection joins cannot multiply balances.
    enriched=[]
    for r in rows:
        outstanding=db().execute('SELECT COALESCE(SUM(balance),0) v FROM accounts a JOIN companies c ON c.id=a.company_id WHERE c.name=?'+(' AND a.assigned_to=?' if g.user['role'] != 'admin' else ''),(r['company_name'],g.user['id']) if g.user['role'] != 'admin' else (r['company_name'],)).fetchone()['v']
        enriched.append(dict(r, outstanding=outstanding, commission=r['collected']*r['commission_rate']/100.0))
    return render_template('reports.html',rows=enriched,start=start,end=end)
@app.route('/users',methods=['GET','POST'])
@admin_required
def users():
    conn=db()
    if request.method=='POST':
        username=request.form.get('username','').strip(); password=request.form.get('password',''); role=request.form.get('role','employee')
        if len(username)<3 or len(password)<12: flash('Username must be at least 3 characters and password at least 12 characters.','error')
        else:
            try: conn.execute('INSERT INTO users(username,password_hash,role,created_at) VALUES(?,?,?,?)',(username,generate_password_hash(password),role,now())); conn.commit(); audit('user_created',username); flash('User created.','success')
            except sqlite3.IntegrityError: flash('That username already exists.','error')
        return redirect(url_for('users'))
    return render_template('users.html',users=conn.execute('SELECT id,username,role,active,created_at FROM users ORDER BY username').fetchall())
@app.route('/users/<int:user_id>/toggle',methods=['POST'])
@admin_required
def toggle_user(user_id):
    if user_id==session['user_id']: flash('You cannot deactivate your own account.','error')
    else:
        db().execute('UPDATE users SET active=CASE active WHEN 1 THEN 0 ELSE 1 END WHERE id=?',(user_id,)); db().commit(); audit('user_status_changed',str(user_id)); flash('User status updated.','success')
    return redirect(url_for('users'))
@app.route('/backup')
@admin_required
def backup():
    dest=BACKUP_DIR / f'collections-backup-{datetime.now().strftime("%Y%m%d-%H%M%S")}.db'
    src=sqlite3.connect(DB_PATH); out=sqlite3.connect(dest); src.backup(out); out.close(); src.close(); audit('backup_created',dest.name)
    return send_file(dest,as_attachment=True,download_name=dest.name,mimetype='application/octet-stream')
@app.route('/export.csv')
@login_required
def export_csv():
    if g.user['role'] == 'admin':
        rows=db().execute('SELECT c.name AS company,a.account_number,a.customer_name,a.email,a.phone,a.address,a.original_balance,a.balance,a.status,a.next_followup FROM accounts a JOIN companies c ON c.id=a.company_id ORDER BY c.name,a.customer_name').fetchall()
    else:
        rows=db().execute('SELECT c.name AS company,a.account_number,a.customer_name,a.email,a.phone,a.address,a.original_balance,a.balance,a.status,a.next_followup FROM accounts a JOIN companies c ON c.id=a.company_id WHERE a.assigned_to=? ORDER BY c.name,a.customer_name',(g.user['id'],)).fetchall()
    out=io.StringIO(); writer=csv.writer(out); writer.writerow(rows[0].keys() if rows else ['company','account_number','customer_name','email','phone','address','original_balance','balance','status','next_followup'])
    for r in rows: writer.writerow(list(r))
    return send_file(io.BytesIO(out.getvalue().encode('utf-8-sig')),as_attachment=True,download_name='collections-accounts.csv',mimetype='text/csv')

@app.context_processor
def common(): return {'current_user': getattr(g,'user',None)}

if __name__ == '__main__':
    init_db()
    # Local development only. Use a production WSGI server and HTTPS in deployment.
    app.run(host='127.0.0.1', port=int(os.environ.get('PORT','5000')), debug=False)
